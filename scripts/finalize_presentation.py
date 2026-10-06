from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import sys, shutil

src=sys.argv[1] if len(sys.argv)>1 else 'presentation/Mav_Metrics_Final.pptx'
tmp=src + '.tmp.pptx'
prs=Presentation(src)

C={
 'bg':'04111F','panel':'081D31','panel2':'0B2741','ink':'F7FAFC','muted':'A9BAC9',
 'blue':'00A3FF','blue2':'2477FF','cyan':'40D9FF','teal':'35D0BA','gold':'F4C95D',
 'line':'193B57','white':'FFFFFF','purple':'9B8CFF'
}
def rgb(v): return RGBColor.from_string(v)
def add_text(slide,txt,x,y,w,h,size=13,color='F7FAFC',bold=False,align=PP_ALIGN.LEFT,font='Aptos',valign=MSO_ANCHOR.MIDDLE):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=txt; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return tb
def add_box(slide,x,y,w,h,fill='081D31',line='193B57'):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.line.color.rgb=rgb(line); shp.line.width=Pt(.9)
def add_pill(slide,txt,x,y,w,color='00A3FF'):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.34))
    shp.fill.solid(); shp.fill.fore_color.rgb=rgb('0D344F'); shp.line.color.rgb=rgb(color); shp.line.width=Pt(.7)
    add_text(slide,txt,x+.06,y+.03,w-.12,.25,7.2,'F7FAFC',True,PP_ALIGN.CENTER)
def top_rule(slide):
    shp=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.055))
    shp.fill.solid(); shp.fill.fore_color.rgb=rgb(C['blue']); shp.line.fill.background()
def footer(slide,n,source=''):
    ln=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.62),Inches(7.06),Inches(12.1),Pt(.6))
    ln.fill.solid(); ln.fill.fore_color.rgb=rgb(C['line']); ln.line.fill.background()
    add_text(slide,f'MAV METRICS  •  {n:02d}',.65,7.10,2,.17,5.8,'6F879A',True)
    if source: add_text(slide,source,2.6,7.08,10.05,.20,5.1,'63798B',False,PP_ALIGN.RIGHT)
def clear_slide(slide):
    sp=slide.shapes._spTree
    for shape in list(slide.shapes): sp.remove(shape._element)

# Slide 5: clean, non-overlapping model pipeline.
s=prs.slides[4]; clear_slide(s); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(C['bg']); top_rule(s)
add_text(s,'MODEL',.68,.35,3.2,.22,8.5,C['blue'],True)
add_text(s,'Simple enough to defend. Useful enough to act on.',.68,.67,11.8,.62,27,C['ink'],True,font='Aptos Display')
add_text(s,'Normalize → score → gap → action. Every step is visible and testable.',.69,1.34,11.7,.45,12.5,C['muted'])
add_box(s,.78,2.05,3.55,3.95,C['panel'],C['line'])
add_text(s,'01  NORMALIZE',1.07,2.34,2.25,.26,8,C['blue'],True)
add_text(s,'Raw counts are not comparable.',1.07,2.88,2.85,.42,15,C['white'],True)
add_text(s,'For followers, views, and mentions:',1.07,3.47,2.75,.30,9.5,C['muted'])
add_text(s,'log(1 + x)',1.07,3.90,2.4,.42,22,C['cyan'],True)
add_text(s,'then percentile-rank every feature from 0–100.',1.07,4.45,2.72,.62,10,C['muted'])
add_pill(s,'SAME SCALE',1.07,5.34,1.25,C['blue'])
add_box(s,4.55,2.05,4,3.95,C['panel2'],C['line'])
add_text(s,'02  SCORE',4.84,2.34,2.25,.26,8,C['teal'],True)
add_text(s,'Marketability score',4.84,2.82,2.9,.35,17,C['white'],True)
weights=[('25%','Performance',C['blue']),('20%','Reach',C['cyan']),('20%','Attention',C['teal']),('15%','Engagement',C['gold']),('10%','Momentum',C['blue2']),('10%','Brand',C['purple'])]
for i,(wt,label,col) in enumerate(weights):
    yy=3.38+i*.36
    dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(4.86),Inches(yy+.07),Inches(.10),Inches(.10))
    dot.fill.solid(); dot.fill.fore_color.rgb=rgb(col); dot.line.fill.background()
    add_text(s,wt,5.08,yy,.58,.23,8.5,col,True); add_text(s,label,5.72,yy,1.9,.23,8.5,C['ink'],True)
add_text(s,'Weights are transparent—and stress-tested next.',4.84,5.67,3.1,.24,8.3,C['muted'])
add_box(s,8.78,2.05,3.77,3.95,C['panel'],C['line'])
add_text(s,'03  FIND THE GAP',9.08,2.34,2.3,.26,8,C['gold'],True)
add_text(s,'Expected = f(Performance)',9.08,2.92,2.95,.38,15,C['white'],True)
add_text(s,'Gap = Marketability − Expected',9.08,3.46,3.02,.44,17,C['white'],True)
ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(9.08),Inches(4.12),Inches(2.92),Pt(.7))
ln.fill.solid(); ln.fill.fore_color.rgb=rgb(C['line']); ln.line.fill.background()
add_text(s,'GAP < 0',9.08,4.38,.78,.30,14,C['gold'],True); add_text(s,'Underexposed relative to performance.',9.98,4.35,2.12,.44,9.3,C['ink'],True)
add_text(s,'GAP > 0',9.08,4.98,.78,.30,14,C['teal'],True); add_text(s,'Brand demand is ahead of the baseline.',9.98,4.95,2.12,.44,9.3,C['ink'],True)
add_pill(s,'SCOUT NEGATIVE GAP + MOMENTUM',9.08,5.51,2.72,C['gold'])
footer(s,5,'Framework formula — weights disclosed and sensitivity-tested')

# Slide 8: a clearly labeled robustness rule, not fake result bars.
s=prs.slides[7]; clear_slide(s); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(C['bg']); top_rule(s)
add_text(s,'ROBUSTNESS',.68,.35,3.2,.22,8.5,C['blue'],True)
add_text(s,'A good recommendation should survive a reasonable argument about weights.',.68,.67,11.8,.62,24.5,C['ink'],True,font='Aptos Display')
add_text(s,'We predefine the robustness test before we look at the winner.',.69,1.34,11.7,.45,12.5,C['muted'])
add_box(s,.78,2.05,7.65,4.22,C['panel'],C['line']); add_text(s,'SENSITIVITY TEST',1.08,2.36,2.3,.27,8,C['blue'],True)
steps=[('01','PERTURB','Move each pillar weight up or down by as much as 10 points.',C['blue']),('02','RENORMALIZE','Scale the new weights back to 100% so scenarios stay comparable.',C['teal']),('03','RE-RANK','Recalculate the opportunity list across the plausible scenarios.',C['gold'])]
for i,(num,label,desc,col) in enumerate(steps):
    yy=2.95+i*.95
    add_text(s,num,1.08,yy,.45,.34,16,col,True); add_text(s,label,1.72,yy,1.45,.30,8.5,col,True); add_text(s,desc,3.16,yy-.02,4.66,.46,10.5,C['ink'],True)
    if i<2:
        ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.16),Inches(yy+.52),Pt(.7),Inches(.33)); ln.fill.solid(); ln.fill.fore_color.rgb=rgb(C['line']); ln.line.fill.background()
add_text(s,'The goal is not to prove one exact coefficient is “right.” It is to see whether the decision survives.',1.08,5.77,6.78,.32,9.5,C['muted'])
add_box(s,8.70,2.05,3.85,4.22,C['panel2'],C['line']); add_text(s,'DECISION RULE',9.03,2.36,2.1,.27,8,C['teal'],True)
add_text(s,'≥ 80%',9.03,3.03,2.4,.62,30,C['teal'],True)
add_text(s,'Keep a candidate only if it remains in the top opportunity set in at least 80% of plausible weight scenarios.',9.03,3.72,2.95,1.12,13,C['white'],True)
add_text(s,'Threshold is a policy choice—not something tuned after seeing the result.',9.03,5.06,2.90,.64,9.5,C['muted'])
add_pill(s,'ROBUST > PRECISE',9.03,5.72,1.70,C['teal'])
footer(s,8,'Robustness rule proposed for model governance; threshold can be adjusted')

# Correct slide numbering in all footers.
for idx,slide in enumerate(prs.slides,1):
    for shape in slide.shapes:
        if shape.has_text_frame and (shape.text or '').strip().startswith('MAV METRICS') and '•' in (shape.text or ''):
            p=shape.text_frame.paragraphs[0]
            if p.runs: p.runs[0].text=f'MAV METRICS  •  {idx:02d}'
            else: p.text=f'MAV METRICS  •  {idx:02d}'
            for run in p.runs:
                run.font.name='Aptos'; run.font.size=Pt(5.8); run.font.bold=True; run.font.color.rgb=rgb('6F879A')

notes={
1:("Joshua","~35 sec","Most marketability models answer the easy question: who is already famous? We wanted to answer the more useful question: who is about to become commercially valuable before everyone else sees it? Mav Metrics combines performance, reach, attention, engagement, momentum, and brand signals into one explainable system. The output is not just a ranking—it is a decision tool.","HANDOFF: I’ll start with the opportunity signal that makes the model useful."),
2:("Joshua","~35 sec","The centerpiece is the marketability gap. We estimate how marketable a player should be from basketball performance, then compare that benchmark with the marketability signals we actually observe. A negative gap does not mean a player is bad at marketing. It means performance may be ahead of commercial attention, which is exactly where a team, agent, or sponsor can act early.","HANDOFF: Shresta will show who uses this and what goes into the score."),
3:("Shresta","~30 sec","We designed the output around decisions, not math for its own sake. A team can use it to decide where to spend content resources. An agent can support a player’s commercial story with evidence. A sponsor can identify rising talent before pricing fully catches up. That is why every score needs both an explanation and a next action.","TRANSITION: To make those decisions defensible, the model uses six public-data pillars."),
4:("Shresta","~35 sec","The six pillars are performance, reach, attention, engagement, momentum, and brand. Performance creates credibility. Reach measures distribution. Attention captures demand beyond a player’s own accounts. Engagement shows audience quality. Momentum measures change. Brand captures commercial proof and fit. We put unlike variables on the same 0-to-100 percentile scale before weighting them.","HANDOFF: Veer will show exactly how those six pillars turn into the gap."),
5:("Veer","~40 sec","The pipeline has three transparent steps. First, we normalize large count variables with a log transform and convert features to percentiles. Second, we combine the six pillars using the disclosed weights. Third, we estimate expected marketability from basketball performance and subtract it from the observed score. That residual is the marketability gap. Negative gap plus strong momentum is where we scout first.","TRANSITION: Then we test the framework against outcomes we did not invent."),
6:("Veer","~35 sec","NBA Communications gives us an outside measure of real fan attention. In the 2024–25 regular season, LeBron James generated 3.23 billion NBA social and digital views, Stephen Curry 2.56 billion, and Luka Dončić 1.82 billion. These are observed outcomes, not our score. A useful model should explain or anticipate behavior like this rather than simply reproduce an internal ranking.","SOURCE: NBA Communications, Apr. 14, 2025."),
7:("Veer","~35 sec","The second outside check is merchandise. For the 2025–26 regular season, Curry ranked first in jersey sales, Luka second, Brunson third, Wembanyama fourth, LeBron fifth, and Dallas rookie Cooper Flagg ninth. We keep jersey sales out of the score and use them as an independent commercial outcome. With enough history, we would evaluate rank correlation, top-K hit rate, and the biggest misses.","TRANSITION: We also test whether our answer survives reasonable changes to the weights."),
8:("Veer","~30 sec","Instead of arguing that one exact set of weights is perfect, we stress-test the decision. We move each pillar by up to ten points, renormalize, and rerank the players. Our governance rule is simple: keep a recommendation only if it stays in the top opportunity set in at least 80 percent of plausible scenarios. The threshold can change; the important part is defining it before seeing the winner.","HANDOFF: Sejal will show how this becomes a reusable product rather than a one-time deck."),
9:("Sejal","~30 sec","The dashboard is deliberately one screen. We can search players, compare them, change assumptions, and see the opportunity map without rebuilding the analysis. The point is not to replace the presentation with software. It is to prove the framework is reusable after the case ends. If we demo it live, we would show only one weight change and one player filter—about thirty seconds total.","TRANSITION: Dallas already gives us a strong example of why momentum belongs in the model."),
10:("Sejal","~35 sec","Cooper Flagg shows the sequence we want the momentum signal to capture. He debuted at number eleven in the midseason jersey rankings, finished ninth for the full 2025–26 regular season as the only rookie in the top fifteen, and then won Rookie of the Year after leading rookies at 21 points per game. The lesson is not that our model discovered Flagg. It is that attention growth can become measurable commercial proof—and we want to detect that pattern earlier in less obvious players.","TRANSITION: That leads to how Dallas could operationalize the framework."),
11:("Sejal","~30 sec","Our recommendation is a quarterly scouting process, not a one-time score. Refresh the same public signals on a fixed date, validate against observable outcomes, create a watchlist, run one small content or sponsor activation, and measure the lift. Start with 10 to 20 players so the team can learn quickly before scaling or adding internal Mavericks data.","HANDOFF: Shresta will close the model section with what we would—and would not—claim from public data."),
12:("Shresta","~30 sec","There are important limits. Public data cannot reveal the exact economics of private endorsement deals, and correlation does not prove causation. We would use Mav Metrics to prioritize where humans should investigate, not to automate decisions. Every recommendation still needs brand-safety review, business judgment, and internal data where available.","HANDOFF: Joshua will close with the one idea we want you to remember."),
13:("Joshua","~20 sec","Mav Metrics is built around one idea: commercial opportunity lives in the gap between what a player is doing on the court and how much value the market has captured off it. We measure the gap, track momentum, validate against real fan behavior, and turn the result into an action. Thank you—we’re happy to take questions.","Q&A ROUTING: Veer = model, weights, validation. Shresta = market logic and limitations. Sejal = dashboard and operating process. Joshua = strategy and close.")
}
for idx,slide in enumerate(prs.slides,1):
    presenter,timing,say,extra=notes[idx]
    tf=slide.notes_slide.notes_text_frame; tf.clear()
    tf.paragraphs[0].text=f'PRESENTER: {presenter}
TIME: {timing}

SAY:
{say}

{extra}'

prs.save(tmp); shutil.move(tmp,src); print(f'Finalized {src}')
