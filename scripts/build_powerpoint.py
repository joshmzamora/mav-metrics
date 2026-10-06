"""Create a simple starter PowerPoint deck.

Run from the repo root:
    python scripts/build_powerpoint.py

This script intentionally keeps the generated deck simple. Polish the final file
manually after replacing demo data with real sourced data.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "generated" / "mav_metrics_deck.pptx"
DATA = ROOT / "data" / "processed" / "player_scores.csv"
FIGURES = ROOT / "figures"
ASSETS = ROOT / "presentation" / "assets"

NAVY = RGBColor(7, 26, 47)
BLUE = RGBColor(0, 125, 197)
GRAY = RGBColor(75, 85, 99)
WHITE = RGBColor(255, 255, 255)


def add_title(slide, title: str, kicker: str = ""):
    if kicker:
        t = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(4.5), Inches(0.25))
        p = t.text_frame.paragraphs[0]
        p.text = kicker.upper()
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = BLUE
    t = slide.shapes.add_textbox(Inches(0.6), Inches(0.7), Inches(11.5), Inches(0.55))
    p = t.text_frame.paragraphs[0]
    p.text = title
    p.runs[0].font.size = Pt(28)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = NAVY


def add_body(slide, text: str, x: float, y: float, w: float, h: float, size: int = 16):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.word_wrap = True
    p = t.text_frame.paragraphs[0]
    p.text = text
    p.runs[0].font.size = Pt(size)
    p.runs[0].font.color.rgb = GRAY


def add_image_if_exists(slide, path: Path, x: float, y: float, w: float):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
        return True
    return False


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides = [
        ("Mav Metrics", "A public-data model for player marketability, momentum, and untapped brand value."),
        ("The question is not who is famous", "Which players are commercially underpriced relative to their performance, attention, and momentum?"),
        ("Public data can still tell a strong story", "Performance, reach, engagement, attention, and commercial proof become comparable after log-scaling and percentile conversion."),
        ("The model has three outputs", "Current marketability, momentum, and marketability gap."),
        ("The map reveals opportunity zones", "Use the performance vs. marketability scatterplot to find underexposed value."),
        ("A good model should match real behavior", "Validate against official jersey-sales lists and NBA digital/social view leaders."),
        ("Starter findings", "Replace demo values with final sourced data before presenting."),
        ("The dashboard makes the model reusable", "The explorer supports Q&A, sensitivity testing, and player comparison."),
        ("How Dallas can use this", "Create a quarterly marketability watchlist for roster and acquisition decisions."),
        ("Limitations", "Public data is incomplete, but the model is transparent and directionally useful."),
        ("Close", "Basketball value creates the stage. Marketability decides who captures it."),
    ]

    for i, (title, body) in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        add_title(slide, title, "Mav Metrics")
        add_body(slide, body, 0.8, 1.7, 5.6, 1.1, 18)
        if i in {4, 6}:
            fig = FIGURES / ("marketability_gap_map.png" if i == 4 else "gap_bar.png")
            add_image_if_exists(slide, fig, 6.8, 1.6, 5.4)
        if i == 0:
            add_image_if_exists(slide, ASSETS / "mav_metrics_logo.png", 7.4, 1.4, 4.2)
        f = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(11.8), Inches(0.25))
        p = f.text_frame.paragraphs[0]
        p.text = f"Mav Metrics | {i+1:02d}"
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = GRAY

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    make_deck()
